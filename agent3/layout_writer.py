"""Generador de Excel de salida en formato LAYOUT."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from agent3.models import COLUMNS_ORDER, TablaIntermediateRow


HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Calibri", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def escribir_tabla_intermedia(
    filas: list[TablaIntermediateRow],
    output_path: Path,
    sheet_name: str = "Tabla_Intermedia",
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for col_idx, col_name in enumerate(COLUMNS_ORDER, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER

    for row_idx, fila in enumerate(filas, 2):
        d = fila.to_dict()
        for col_idx, col_name in enumerate(COLUMNS_ORDER, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=d[col_name])
            cell.font = DATA_FONT

    for col_idx, col_name in enumerate(COLUMNS_ORDER, 1):
        max_len = len(col_name)
        for row_idx in range(2, len(filas) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 30)

    ws.auto_filter.ref = ws.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))

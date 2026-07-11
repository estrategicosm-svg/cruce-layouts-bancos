"""Genera Excel usando los templates exactos y llenando CRUCE BANCARIO."""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from openpyxl import load_workbook

from cruce_layouts_bancos.models import MovimientoBanco

ASSETS = Path(__file__).parent / "assets"


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def generar_excel(
    filas_egresos: list,
    filas_ingresos: list,
    movimientos: list[MovimientoBanco],
    cruce_map: dict[str, str],
    archivo_salida: str,
) -> dict:
    Path(archivo_salida).parent.mkdir(parents=True, exist_ok=True)

    # 1. Copy egresos template as base
    shutil.copy2(ASSETS / "LAYOUT_CARGA_EGRESOS_OUTPUT.xlsx", archivo_salida)
    wb = load_workbook(archivo_salida)
    ws_eg = wb["DATOS"]

    # Build index: (empresa_upper, poliza, uuid) -> grupo_id
    idx_eg = {}
    for f in filas_egresos:
        key = (f.empresa.upper().strip(), f.poliza.strip(), f.uuid.strip())
        idx_eg[key] = f"{f.empresa}|{f.poliza}"

    col_cruce = ws_eg.max_column  # col 12 = CRUCE BANCARIO
    for row_idx in range(3, ws_eg.max_row + 1):
        emp = str(ws_eg.cell(row_idx, 1).value or "").strip()
        pol = str(ws_eg.cell(row_idx, 2).value or "").strip()
        uuid_val = str(ws_eg.cell(row_idx, 3).value or "").strip()
        grupo_id = idx_eg.get((emp.upper(), pol, uuid_val))
        if grupo_id and grupo_id in cruce_map:
            ws_eg.cell(row_idx, col_cruce).value = cruce_map[grupo_id]

    # 2. Add ingresos sheet from template
    wb_ing_src = load_workbook(ASSETS / "LAYOUT_CEDULA_INGRESOS_OUTPUT.xlsx")
    ws_ing_src = wb_ing_src["DATOS"]
    ws_ing = wb.create_sheet("INGRESOS")
    for row_idx in range(1, ws_ing_src.max_row + 1):
        for col_idx in range(1, ws_ing_src.max_column + 1):
            ws_ing.cell(row_idx, col_idx).value = ws_ing_src.cell(row_idx, col_idx).value
    wb_ing_src.close()

    idx_ing = {}
    for f in filas_ingresos:
        key = (f.empresa.upper().strip(), f.poliza.strip(), f.uuid.strip())
        idx_ing[key] = f"{f.empresa}|{f.poliza}"

    col_cruce_ing = ws_ing.max_column
    for row_idx in range(3, ws_ing.max_row + 1):
        emp = str(ws_ing.cell(row_idx, 1).value or "").strip()
        pol = str(ws_ing.cell(row_idx, 2).value or "").strip()
        uuid_val = str(ws_ing.cell(row_idx, 3).value or "").strip()
        grupo_id = idx_ing.get((emp.upper(), pol, uuid_val))
        if grupo_id and grupo_id in cruce_map:
            ws_ing.cell(row_idx, col_cruce_ing).value = cruce_map[grupo_id]

    # 3. Add movimientos sheet from template
    wb_mov_src = load_workbook(ASSETS / "MOVIMIENTOS_BANCARIOS_OUTPUT.xlsx")
    ws_mov_src = wb_mov_src["ESTADOS_CUENTA"]
    ws_mov = wb.create_sheet("MOVIMIENTOS_BANCARIOS")
    for row_idx in range(1, ws_mov_src.max_row + 1):
        for col_idx in range(1, ws_mov_src.max_column + 1):
            ws_mov.cell(row_idx, col_idx).value = ws_mov_src.cell(row_idx, col_idx).value
    wb_mov_src.close()

    # Fill CRUCE BANCARIO + POLIZA in movimientos
    col_cruce_cargo = 15
    col_cruce_abono = 17
    col_poliza = 20
    mov_by_archivo = {}
    for m in movimientos:
        mov_by_archivo.setdefault(m.archivo_origen, []).append(m)

    for row_idx in range(2, ws_mov.max_row + 1):
        archivo = str(ws_mov.cell(row_idx, 1).value or "")
        cargo_val = ws_mov.cell(row_idx, 14).value
        abono_val = ws_mov.cell(row_idx, 16).value
        fname = Path(archivo).name

        for m in mov_by_archivo.get(fname, []):
            matched = False
            if m.cruce_cargo and cargo_val:
                try:
                    if abs(float(cargo_val) - float(m.cargo)) < 0.01:
                        ws_mov.cell(row_idx, col_cruce_cargo).value = m.cruce_cargo
                        if m.poliza_relacionada:
                            ws_mov.cell(row_idx, col_poliza).value = m.poliza_relacionada
                        matched = True
                except (ValueError, TypeError):
                    pass
            if not matched and m.cruce_abono and abono_val:
                try:
                    if abs(float(abono_val) - float(m.abono)) < 0.01:
                        ws_mov.cell(row_idx, col_cruce_abono).value = m.cruce_abono
                        if m.poliza_relacionada:
                            ws_mov.cell(row_idx, col_poliza).value = m.poliza_relacionada
                except (ValueError, TypeError):
                    pass
            if matched:
                break

    # Remove default empty sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(archivo_salida)
    raw = Path(archivo_salida).read_bytes()
    return {
        "archivo": archivo_salida,
        "tamaño": len(raw),
        "sha256": _sha256(raw),
        "hojas": wb.sheetnames,
    }

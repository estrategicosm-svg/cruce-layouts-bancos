"""Dashboard Streamlit - Cruce de layouts vs estados de cuenta bancarios."""
from __future__ import annotations


def main():
    import importlib.util
    import io
    import sys
    import zipfile
    from datetime import datetime, timezone
    from pathlib import Path

    _HERE = Path(__file__).resolve().parent
    _REPO = _HERE.parent

    if str(_REPO) not in sys.path:
        sys.path.insert(0, str(_REPO))

    spec = importlib.util.spec_from_file_location(
        "cruce_layouts_bancos.app", str(_HERE / "app.py")
    )
    _mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(_mod)
    ejecutar_cruce = _mod.ejecutar_cruce

    import pandas as pd
    import streamlit as st
    from openpyxl import load_workbook

    st.set_page_config(
        page_title="Conciliador Layouts vs Bancos",
        page_icon="",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    st.markdown("""
    <style>
        .stMetric > div { padding: 12px 16px; border-radius: 8px; }
        div[data-testid="stMetricValue"] { font-size: 1.6rem !important; }
        div[data-testid="stMetricDelta"] { font-size: 0.85rem !important; }
        .block-container { padding-top: 1.5rem; }
        section[data-testid="stSidebar"] > div { padding-top: 1rem; }
    </style>
    """, unsafe_allow_html=True)

    def _combinar_bancos_a_zip(
        zip_file, pdf_files: list
    ) -> tuple[bytes, str, int, int]:
        buf = io.BytesIO()
        n_pdfs = 0
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            if zip_file is not None:
                zf.writestr(zip_file.name, zip_file.getvalue())
            if pdf_files:
                for pf in pdf_files:
                    zf.writestr(f"PDF_SUELTOS/{pf.name}", pf.getvalue())
                    n_pdfs += 1
        buf.seek(0)
        nombre = "BANCOS_COMBINADOS.zip"
        n_zip = 1 if zip_file else 0
        return buf.getvalue(), nombre, n_zip, n_pdfs

    def _leer_hoja_excel(excel_bytes: bytes, sheet_name: str, min_row: int = 2):
        wb = load_workbook(io.BytesIO(excel_bytes), read_only=True)
        try:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(min_row=min_row, values_only=True))
        finally:
            wb.close()
        if not rows:
            return [], []
        headers = [str(h or f"col_{i}") for i, h in enumerate(rows[0])]
        return headers, rows[1:]

    # ── Sidebar ──────────────────────────────────────────────────────────────
    with st.sidebar:
        st.markdown("## Configuracion")

        st.markdown("### Archivos de entrada")
        f_egresos = st.file_uploader(
            "Layout Egresos (.xlsx)", type=["xlsx"],
            help="Layout de carga de egresos del sistema contable",
        )
        f_ingresos = st.file_uploader(
            "Layout Ingresos (.xlsx)", type=["xlsx"],
            help="Layout de cedula de ingresos del sistema contable",
        )
        f_bancos_zip = st.file_uploader(
            "ZIP Estados de cuenta (.zip)", type=["zip"],
            help="ZIP con PDFs y/o Excel de estados de cuenta bancarios",
        )
        f_bancos_pdfs = st.file_uploader(
            "PDFs sueltos de estados de cuenta", type=["pdf"],
            accept_multiple_files=True,
            help="Archivos PDF individuales de estados de cuenta bancarios",
        )

        st.markdown("---")
        st.markdown("### Tolerancia de cruce")
        col_t1, col_t2 = st.columns(2)
        tol_mxn = col_t1.number_input(
            "MXN ($)", min_value=0.0, max_value=10000.0, value=1.0, step=0.5,
            help="Tolerancia maxima en pesos para aceptar un cruce",
        )
        tol_usd = col_t2.number_input(
            "USD ($)", min_value=0.0, max_value=1000.0, value=0.10, step=0.01,
            help="Tolerancia maxima en dolares para aceptar un cruce",
        )

        st.markdown("---")
        st.markdown("### Descargas")
        assets = _HERE / "assets"
        for idx, (fname, label) in enumerate([
            ("LAYOUT_CARGA_EGRESOS_OUTPUT.xlsx", "Template Egresos"),
            ("LAYOUT_CEDULA_INGRESOS_OUTPUT.xlsx", "Template Ingresos"),
        ]):
            p = assets / fname
            if p.exists():
                st.download_button(
                    label=label,
                    data=p.read_bytes(),
                    file_name=fname,
                    type="secondary",
                    use_container_width=True,
                    key=f"tpl_{idx}",
                )

        st.markdown("---")
        has_input = f_egresos and f_ingresos and (f_bancos_zip or f_bancos_pdfs)
        ejecutar = st.button(
            "Ejecutar cruce",
            type="primary",
            use_container_width=True,
            disabled=not has_input,
        )

        if st.session_state.get("excel_resultado"):
            st.markdown("---")
            st.download_button(
                label="Descargar Excel final",
                data=st.session_state["excel_resultado"],
                file_name=st.session_state.get("excel_nombre", "CRUCE_LAYOUTS_VS_BANCOS.xlsx"),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
                key="sidebar_download",
            )

    # ── Main ─────────────────────────────────────────────────────────────────
    st.markdown("## Conciliador Layouts vs Bancos")

    if ejecutar:
        zip_bytes_data, nombre_zip, n_zip, n_pdfs = _combinar_bancos_a_zip(
            f_bancos_zip, f_bancos_pdfs
        )

        with st.spinner("Procesando..."):
            try:
                result = ejecutar_cruce(
                    egresos_bytes=f_egresos.getvalue(),
                    nombre_egresos=f_egresos.name,
                    ingresos_bytes=f_ingresos.getvalue(),
                    nombre_ingresos=f_ingresos.name,
                    zip_bancos_bytes=zip_bytes_data,
                    nombre_zip=nombre_zip,
                    tolerancia_mxn=tol_mxn,
                    tolerancia_usd=tol_usd,
                )
            except Exception as exc:
                st.error(f"Error durante el cruce: {exc}")
                st.stop()

        if result.get("EXCEL_ERROR"):
            st.error(f"El cruce se ejecuto pero fallo la generacion del Excel: {result['EXCEL_ERROR']}")
            st.stop()

        excel_bytes_check = result.get("EXCEL_BYTES")
        if not isinstance(excel_bytes_check, (bytes, bytearray)):
            st.error(
                "El motor no devolvio un Excel valido. "
                f"Tipo recibido: {type(excel_bytes_check).__name__}"
            )
            st.stop()

        excel_bytes_check = bytes(excel_bytes_check)
        if not excel_bytes_check.startswith(b"PK"):
            st.error("El archivo generado no tiene formato XLSX valido.")
            st.stop()

        st.session_state["cruce_result"] = result
        st.session_state["excel_resultado"] = excel_bytes_check
        st.session_state["excel_nombre"] = result.get("EXCEL_NOMBRE", "CRUCE_LAYOUTS_VS_BANCOS.xlsx")
        st.session_state["excel_metadata"] = {
            "tamaño": result.get("EXCEL_TAMANO", len(excel_bytes_check)),
            "sha256": result.get("EXCEL_SHA256", ""),
            "hojas": result.get("EXCEL_HOJAS", []),
            "fecha_generacion": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
            "tolerancia_mxn": result.get("tolerancia_mxn", 1.0),
            "tolerancia_usd": result.get("tolerancia_usd", 0.10),
        }

    if not st.session_state.get("excel_resultado"):
        st.info("Sube los archivos en la barra lateral y presiona **Ejecutar cruce**.")
        st.stop()

    result = st.session_state.get("cruce_result", {})
    excel_bytes = st.session_state["excel_resultado"]

    # ── Alertas de lectura ────────────────────────────────────────────────────
    if ejecutar:
        if result.get("ERRORES_ARCHIVOS"):
            st.error(
                f"**{len(result['ERRORES_ARCHIVOS'])} archivo(s) fallaron al procesarse:**"
            )
            for err in result["ERRORES_ARCHIVOS"]:
                st.markdown(f"- **{err.archivo}** ({err.tipo}): {err.error}")

        for adv in result.get("ADVERTENCIAS_LECTURA", []):
            st.warning(adv)

    # ── Boton de descarga principal ──────────────────────────────────────────
    meta = st.session_state.get("excel_metadata", {})
    st.success(
        f"Excel generado: **{st.session_state.get('excel_nombre', 'CRUCE_LAYOUTS_VS_BANCOS.xlsx')}** "
        f"| {meta.get('tamaño', 0):,} bytes "
        f"| Hojas: {', '.join(meta.get('hojas', []))}"
    )

    col_dl1, col_dl2 = st.columns([1, 3])
    with col_dl1:
        st.download_button(
            label="Descargar Excel final",
            data=excel_bytes,
            file_name=st.session_state.get("excel_nombre", "CRUCE_LAYOUTS_VS_BANCOS.xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
            key="main_download",
        )

    # ── KPIs ─────────────────────────────────────────────────────────────────
    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Cruces", result.get("CRUCES_ENCONTRADOS_TOTAL", 0))
    c2.metric("Multiples", result.get("MULTIPLES_CANDIDATOS_TOTAL", 0))
    c3.metric("Sin candidato", result.get("SIN_CANDIDATO_TOTAL", 0))
    c4.metric("Movimientos", result.get("MOVIMIENTOS_BANCARIOS", 0))
    c5.metric("TDC procesadas", result.get("MOVIMIENTOS_TDC", 0))
    c6.metric("Tiempo", f"{result.get('TIEMPO', 0):.1f}s")

    # ── Resumen archivos ZIP ─────────────────────────────────────────────────
    with st.expander("Archivos procesados del ZIP", expanded=False):
        st.markdown(f"""
        | Metrica | Cantidad |
        |---------|----------|
        | Archivos totales | {result.get('ARCHIVOS_TOTALES_ZIP', 0)} |
        | Excel procesados OK | {result.get('ARCHIVOS_XLSX_OK', 0)} |
        | Excel fallidos | {result.get('ARCHIVOS_XLSX_FALLIDOS', 0)} |
        | PDFs en ZIP | {result.get('ARCHIVOS_PDF_TOTAL', 0)} |
        | PDFs procesados OK | {result.get('ARCHIVOS_PDF_OK', 0)} |
        | PDFs fallidos | {result.get('ARCHIVOS_PDF_FALLIDOS', 0)} |
        | Tesseract OCR | {"Disponible" if result.get('TESSERACT_DISPONIBLE') else "NO disponible"} |
        """)

    st.markdown("---")

    # ── Tabs ─────────────────────────────────────────────────────────────────
    tab_res, tab_eg, tab_ing, tab_mov, tab_arch = st.tabs([
        "Resumen", "Egresos", "Ingresos", "Movimientos", "Archivo",
    ])

    # ── Resumen ──────────────────────────────────────────────────────────────
    with tab_res:
        st.markdown("### Cuadre por grupos")
        r = result
        st.markdown(f"""
        | Modulo | Encontrados | Multiples | Sin candidato | Mezclados | Total |
        |--------|------------|-----------|---------------|-----------|-------|
        | **Egresos** | {r.get('CRUCES_ENCONTRADOS_EGRESOS', 0)} | {r.get('MULTIPLES_CANDIDATOS_EGRESOS', 0)} | {r.get('SIN_CANDIDATO_EGRESOS', 0)} | — | {r.get('GRUPOS_EGRESOS', 0)} |
        | **Ingresos** | {r.get('CRUCES_ENCONTRADOS_INGRESOS', 0)} | {r.get('MULTIPLES_CANDIDATOS_INGRESOS', 0)} | {r.get('SIN_CANDIDATO_INGRESOS', 0)} | — | {r.get('GRUPOS_INGRESOS', 0)} |
        | **Total** | **{r.get('CRUCES_ENCONTRADOS_TOTAL', 0)}** | **{r.get('MULTIPLES_CANDIDATOS_TOTAL', 0)}** | **{r.get('SIN_CANDIDATO_TOTAL', 0)}** | {r.get('MONEDAS_MIXTAS', 0)} | **{r.get('GRUPOS_EGRESOS', 0) + r.get('GRUPOS_INGRESOS', 0)}** |
        """)

        st.markdown("### Movimientos bancarios")
        st.markdown(f"""
        | Tipo | Cantidad |
        |------|----------|
        | Cargos (egresos) | {r.get('MOVIMIENTOS_CARGOS', 0)} |
        | Abonos (ingresos) | {r.get('MOVIMIENTOS_ABONOS', 0)} |
        | TDC (procesadas) | {r.get('MOVIMIENTOS_TDC', 0)} |
        | **Total** | **{r.get('MOVIMIENTOS_BANCARIOS', 0)}** |
        """)

        st.markdown("### Cruces asignados")
        cm = r.get("cruce_map", {})
        if cm:
            df_cruces = pd.DataFrame([
                {"Grupo": k, "Cruce": v} for k, v in sorted(cm.items())
            ])
            st.dataframe(df_cruces, use_container_width=True, height=300)

    # ── Egresos ──────────────────────────────────────────────────────────────
    with tab_eg:
        st.markdown("### Layout de egresos con CRUCE BANCARIO")
        try:
            headers, data = _leer_hoja_excel(excel_bytes, "EGRESOS")
            df_eg = pd.DataFrame(data, columns=headers) if headers else pd.DataFrame()

            if not df_eg.empty and "CRUCE BANCARIO" in df_eg.columns:
                matched = df_eg["CRUCE BANCARIO"].notna().sum()
                st.caption(f"{matched} de {len(df_eg)} filas con cruce asignado")
                highlight = df_eg[df_eg["CRUCE BANCARIO"].notna()]
                if not highlight.empty:
                    st.dataframe(highlight, use_container_width=True, height=400)
            elif not df_eg.empty:
                st.dataframe(df_eg, use_container_width=True, height=400)
        except Exception as e:
            st.error(f"Error leyendo egresos: {e}")

    # ── Ingresos ─────────────────────────────────────────────────────────────
    with tab_ing:
        st.markdown("### Layout de ingresos con CRUCE BANCARIO")
        try:
            headers, data = _leer_hoja_excel(excel_bytes, "INGRESOS")
            df_ing = pd.DataFrame(data, columns=headers) if headers else pd.DataFrame()

            if not df_ing.empty and "CRUCE BANCARIO" in df_ing.columns:
                matched = df_ing["CRUCE BANCARIO"].notna().sum()
                st.caption(f"{matched} de {len(df_ing)} filas con cruce asignado")
                highlight = df_ing[df_ing["CRUCE BANCARIO"].notna()]
                if not highlight.empty:
                    st.dataframe(highlight, use_container_width=True, height=400)
            elif not df_ing.empty:
                st.dataframe(df_ing, use_container_width=True, height=400)
        except Exception as e:
            st.error(f"Error leyendo ingresos: {e}")

    # ── Movimientos ──────────────────────────────────────────────────────────
    with tab_mov:
        st.markdown("### Movimientos bancarios")
        try:
            wb = load_workbook(io.BytesIO(excel_bytes), read_only=True)
            try:
                ws = wb["MOVIMIENTOS_BANCARIOS"]

                summary_start = None
                for row in range(1, ws.max_row + 1):
                    v = ws.cell(row, 1).value
                    if v and "RESUMEN" in str(v).upper():
                        summary_start = row
                        break

                if summary_start:
                    st.caption(f"Datos: filas 2-{summary_start - 2} | Resumen: filas {summary_start}-{ws.max_row}")

                    data_rows = []
                    for row in range(2, summary_start):
                        vals = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
                        if any(v is not None for v in vals):
                            data_rows.append(vals)

                    hdrs = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
                    df_mov = pd.DataFrame(data_rows, columns=hdrs)

                    cruce_cargo = df_mov["CRUCE BANCARIO"].notna().sum() if "CRUCE BANCARIO" in df_mov.columns else 0
                    poliza_col = "POLIZA" if "POLIZA" in df_mov.columns else None
                    poliza_filled = df_mov[poliza_col].notna().sum() if poliza_col else 0
                    st.caption(f"CRUCE BANCARIO (cargo): {cruce_cargo} | POLIZA asignada: {poliza_filled}")

                    show_cols = [c for c in ["empresa_detectada", "banco_detectada", "cuenta_detectada",
                                              "moneda_detectada", "fecha_movimiento", "descripcion_original",
                                              "cargo", "CRUCE BANCARIO", "abono", "POLIZA"]
                                 if c in df_mov.columns]
                    st.dataframe(df_mov[show_cols], use_container_width=True, height=400)

                    st.markdown("---")
                    st.markdown("#### Resumen por banco/cuenta/moneda")
                    summary_rows = []
                    for row in range(summary_start + 2, ws.max_row + 1):
                        vals = [ws.cell(row, c).value for c in range(1, 9)]
                        if any(v is not None for v in vals):
                            summary_rows.append(vals)

                    if summary_rows:
                        sum_headers = ["Archivo", "Empresa", "Moneda", "Total Cargos",
                                       "Total Abonos", "Neto", "Movimientos", "Validacion"]
                        df_sum = pd.DataFrame(summary_rows, columns=sum_headers)
                        cuadra = df_sum[df_sum["Validacion"] == "CUADRA"]
                        no_cuadra = df_sum[df_sum["Validacion"] != "CUADRA"]
                        st.caption(f"Cuadran: {len(cuadra)} | No cuadran: {len(no_cuadra)}")
                        st.dataframe(df_sum, use_container_width=True, height=350)

                        if not no_cuadra.empty:
                            st.warning(f"{len(no_cuadra)} cuentas no cuadran. Revisa los montos.")
                            st.dataframe(no_cuadra, use_container_width=True)
                else:
                    st.warning("No se encontro seccion de resumen en el archivo.")
            finally:
                wb.close()
        except Exception as e:
            st.error(f"Error leyendo movimientos: {e}")

    # ── Archivo ──────────────────────────────────────────────────────────────
    with tab_arch:
        st.markdown("### Descargar resultado")

        excel_data = st.session_state.get("excel_resultado")
        meta_arch = st.session_state.get("excel_metadata", {})
        excel_name = st.session_state.get("excel_nombre", "CRUCE_LAYOUTS_VS_BANCOS.xlsx")

        if excel_data:
            st.download_button(
                label="Descargar Excel final",
                data=excel_data,
                file_name=excel_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
                key="arch_download",
            )
            st.markdown(f"""
            | Propiedad | Valor |
            |-----------|-------|
            | Archivo | `{excel_name}` |
            | Tamano | {meta_arch.get("tamaño", 0):,} bytes |
            | Fecha generacion | {meta_arch.get("fecha_generacion", "N/A")} |
            | Hojas | {', '.join(meta_arch.get("hojas", []))} |
            | SHA256 | `{meta_arch.get("sha256", "N/A")[:32]}...` |
            | Tolerancia MXN | ${meta_arch.get("tolerancia_mxn", 0):.2f} |
            | Tolerancia USD | ${meta_arch.get("tolerancia_usd", 0):.2f} |
            """)
        else:
            st.warning("No existe un Excel generado. Ejecuta nuevamente el cruce.")


if __name__ == "__main__":
    main()

"""Tests para validacion de descarga de Excel."""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from cruce_layouts_bancos.output_writer import generar_excel
from cruce_layouts_bancos.models import MovimientoBanco, FilaLayout


ASSETS = Path(__file__).resolve().parent.parent / "assets"


def _make_mock_fila(empresa="CSC", poliza="POL001", uuid="UUID1",
                    total_pesos=1000.0, moneda="MXN"):
    from decimal import Decimal
    return FilaLayout(
        empresa=empresa,
        poliza=poliza,
        uuid=uuid,
        rfc="XAXX010101000",
        nombre="TEST",
        fecha_factura="2024-02-15",
        fecha_operacion="2024-02-15",
        moneda=moneda,
        total_pesos=Decimal(str(total_pesos)),
        total_dls=Decimal("0"),
        banco="BANAMEX",
    )


def _make_mock_movimiento(cargo=1000.0, abono=0.0, empresa="CSC",
                          archivo="INTRA_BANAMEX_CTA0688_MXN_FEB2024.xlsx"):
    return MovimientoBanco(
        archivo_origen=archivo,
        ruta_interna_zip=archivo,
        empresa_detectada=empresa,
        banco_detectado="BANAMEX",
        cuenta_detectada="CTA0688",
        cuenta_ultimos_4="0688",
        moneda_detectada="MXN",
        mes_detectado="FEB2024",
        fecha_movimiento="2024-02-15",
        descripcion_original="PAGO PROVEEDOR",
        referencia_usada="REF001",
        clave_rastreo="",
        autorizacion="",
        cargo=cargo,
        abono=abono,
        saldo=0,
        tipo_movimiento="CARGO" if cargo > 0 else "ABONO",
        es_tdc=False,
    )


class TestExcelGeneracionEnMemoria:
    def test_generar_excel_retorna_bytes(self):
        filas_eg = [_make_mock_fila()]
        filas_ing = []
        movs = [_make_mock_movimiento()]
        cruce_map = {"CSC|POL001": "CSC-POL001-001"}

        result = generar_excel(filas_eg, filas_ing, movs, cruce_map)

        assert "bytes" in result
        assert isinstance(result["bytes"], bytes)
        assert len(result["bytes"]) > 0

    def test_bytes_empiezan_con_firma_xlsx(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        raw = result["bytes"]
        assert raw[:2] == b"PK"

    def test_excel_abre_con_openpyxl(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        wb = load_workbook(io.BytesIO(result["bytes"]))
        assert wb is not None
        wb.close()

    def test_tiene_tres_hojas(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        wb = load_workbook(io.BytesIO(result["bytes"]))
        assert "EGRESOS" in wb.sheetnames
        assert "INGRESOS" in wb.sheetnames
        assert "MOVIMIENTOS_BANCARIOS" in wb.sheetnames
        assert len(wb.sheetnames) == 3
        wb.close()

    def test_archivo_salida_none_no_escribe_a_disco(self, tmp_path):
        outfile = str(tmp_path / "test.xlsx")
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
            archivo_salida=None,
        )
        assert not Path(outfile).exists()

    def test_archivo_salida_escribe_a_disco(self, tmp_path):
        outfile = str(tmp_path / "test.xlsx")
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
            archivo_salida=outfile,
        )
        assert Path(outfile).exists()
        assert Path(outfile).read_bytes()[:2] == b"PK"

    def test_bytes_no_vacios(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        assert len(result["bytes"]) > 1000


class TestSessionStatePersistencia:
    def test_session_state_cubre_cambio_de_pestana(self):
        """Simula: ejecutar, guardar en session_state, y leer despues."""
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        session_state = {
            "excel_resultado": result["bytes"],
            "excel_nombre": "CRUCE_LAYOUTS_VS_BANCOS.xlsx",
            "excel_metadata": {
                "tamaño": result["tamaño"],
                "sha256": result["sha256"],
                "hojas": result["hojas"],
            },
        }
        assert session_state["excel_resultado"] is not None
        assert len(session_state["excel_resultado"]) > 0
        assert session_state["excel_nombre"] == "CRUCE_LAYOUTS_VS_BANCOS.xlsx"

    def test_excel_lee_hojas_desde_session_state(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        excel_bytes = result["bytes"]
        wb = load_workbook(io.BytesIO(excel_bytes), read_only=True)
        assert "EGRESOS" in wb.sheetnames
        ws = wb["EGRESOS"]
        assert ws.max_row >= 3
        wb.close()

    def test_metadata_completa(self):
        result = generar_excel(
            [_make_mock_fila()], [], [_make_mock_movimiento()],
            {"CSC|POL001": "CRUCE-001"},
        )
        assert result["tamaño"] > 0
        assert len(result["sha256"]) == 64
        assert "EGRESOS" in result["hojas"]
        assert "INGRESOS" in result["hojas"]
        assert "MOVIMIENTOS_BANCARIOS" in result["hojas"]

"""E2E test with real production data.

Loads:
  - LAYOUT_CARGA_EGRESOS.xlsx
  - LAYOUT_CEDULA_INGRESOS.xlsx
  - ESTADOS_DE_CTA_RENOMBRADO.zip
  - XML_feb_24.zip
  - XML_RECIBIDOS_FEB_24.zip

Produces full evidence for DICTAMEN.
"""
from __future__ import annotations

import hashlib
import io
import os
import sys
import tempfile
import time
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))

from conciliacion.conciliador import Conciliador
from conciliacion.conciliador_ingresos import ConciliadorIngresos
from core.models import CedulaRegistro, CedulaIngresoRegistro
from core.utils import to_decimal, to_datetime, normalizar_uuid, normalizar_referencia
from exportadores.excel_generator import ExcelGenerator
from exportadores.reportes import (
    registros_a_dataframe,
    registros_ingresos_a_dataframe,
    resumen_a_dataframe,
)
from parsers.banco_parser import BancoParser
from parsers.xml_parser import XMLParser

PAQUETE = Path(r"C:\Users\USER\Downloads\PAQUETE PARA SUBIR")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def _to_decimal(v) -> Decimal:
    if v is None:
        return Decimal("0")
    try:
        return Decimal(str(v).replace(",", "").strip())
    except Exception:
        return Decimal("0")


def _to_datetime(v) -> datetime:
    if v is None:
        return datetime(2024, 1, 1)
    if isinstance(v, datetime):
        return v
    return to_datetime(v)


def _load_egresos() -> list[CedulaRegistro]:
    wb = openpyxl.load_workbook(
        str(PAQUETE / "LAYOUT_CARGA_EGRESOS.xlsx"),
        read_only=True, data_only=True,
    )
    ws = wb["DATOS"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_row = rows[1]
    cols = {str(h).strip().upper(): i for i, h in enumerate(header_row) if h}
    egresos = []
    for row in rows[2:]:
        if not row or not any(row):
            continue
        try:
            egresos.append(CedulaRegistro(
                poliza=str(row[cols.get("POLIZA", 1)] or ""),
                cliente=str(row[cols.get("PROVEEDOR", 4)] or ""),
                rfc=str(row[cols.get("RFC_PROVEEDOR", 3)] or ""),
                uuid=normalizar_uuid(row[cols.get("UUID", 2)] or ""),
                importe=_to_decimal(row[cols.get("TOTAL PESOS", 8)]),
                base_iva_16=_to_decimal(row[cols.get("TOTAL PESOS", 8)]),
                base_iva_8=Decimal("0"),
                base_iva_0=Decimal("0"),
                exentos=Decimal("0"),
                iva=Decimal("0"),
                retenciones=Decimal("0"),
                moneda=str(row[cols.get("MONEDA", 7)] or "MXN"),
                tipo_cambio=Decimal("1"),
                fecha_pago=_to_datetime(row[cols.get("FECHA_PAGO", 6)]),
                banco=str(row[cols.get("BANCO", 10)] or ""),
                cruce_bancario="",
            ))
        except Exception:
            continue
    return egresos


def _load_ingresos() -> list[CedulaIngresoRegistro]:
    wb = openpyxl.load_workbook(
        str(PAQUETE / "LAYOUT_CEDULA_INGRESOS.xlsx"),
        read_only=True, data_only=True,
    )
    ws = wb["DATOS"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_row = rows[1]
    cols = {str(h).strip().upper(): i for i, h in enumerate(header_row) if h}
    ingresos = []
    for row in rows[2:]:
        if not row or not any(row):
            continue
        try:
            ingresos.append(CedulaIngresoRegistro(
                uuid=normalizar_uuid(row[cols.get("UUID", 2)] or ""),
                cliente=str(row[cols.get("CLIENTE", 4)] or ""),
                rfc=str(row[cols.get("RFC_CLIENTE", 3)] or ""),
                factura=str(row[cols.get("POLIZA", 1)] or ""),
                fecha=_to_datetime(row[cols.get("FECHA_COBRO", 6)]),
                total=_to_decimal(row[cols.get("TOTAL PESOS", 8)]),
                moneda=str(row[cols.get("MONEDA", 7)] or "MXN"),
                amount_mxn=_to_decimal(row[cols.get("TOTAL PESOS", 8)]),
                amount_usd=_to_decimal(row[cols.get("TOTAL DLS", 9)]),
                forma_pago="",
                folio_transferencia=normalizar_referencia(row[cols.get("BANCO", 10)] or ""),
                descripcion=str(row[cols.get("EMPRESA", 0)] or ""),
            ))
        except Exception:
            continue
    return ingresos


def main() -> None:
    print("=" * 80)
    print("E2E TEST — SAT CONCILIADOR IVA")
    print("=" * 80)
    t0 = time.time()

    # ── 1. Parse cedula egresos ──
    print("\n[1] PARSEANDO CEDULA EGRESOS...")
    cedula = _load_egresos()
    print(f"    FILAS_EGRESOS: {len(cedula)}")

    # ── 2. Parse cedula ingresos ──
    print("\n[2] PARSEANDO CEDULA INGRESOS...")
    cedula_ingresos = _load_ingresos()
    print(f"    FILAS_INGRESOS: {len(cedula_ingresos)}")

    # ── 3. Parse XML emitidos ──
    print("\n[3] PARSEANDO XML EMITIDOS...")
    xml_parser = XMLParser()
    xml_emitidos_file = PAQUETE / "XML_feb_24.zip"
    resultado_emit = xml_parser.parsear_zip(xml_emitidos_file.read_bytes(), origen_xml="EMITIDO")
    cfdis_emitidos = resultado_emit.exitosos
    print(f"    XML_EMITIDOS: {resultado_emit.exitosos_count}")
    print(f"    XML_EMITIDOS_FALLOS: {resultado_emit.fallos_count}")
    if resultado_emit.fallos:
        for fl in resultado_emit.fallos[:5]:
            print(f"      FALLO: {fl['archivo']}: {fl['error'][:80]}")

    # ── 4. Parse XML recibidos ──
    print("\n[4] PARSEANDO XML RECIBIDOS...")
    xml_recibidos_file = PAQUETE / "XML_RECIBIDOS_FEB_24.zip"
    resultado_rec = xml_parser.parsear_zip(xml_recibidos_file.read_bytes(), origen_xml="RECIBIDO")
    cfdis_recibidos = resultado_rec.exitosos
    print(f"    XML_RECIBIDOS: {resultado_rec.exitosos_count}")
    print(f"    XML_RECIBIDOS_FALLOS: {resultado_rec.fallos_count}")
    if resultado_rec.fallos:
        for fl in resultado_rec.fallos[:5]:
            print(f"      FALLO: {fl['archivo']}: {fl['error'][:80]}")

    cfdis = cfdis_emitidos + cfdis_recibidos
    print(f"    XML_TOTALES: {len(cfdis)}")

    # ── 5. Parse banco (ZIP bancario) ──
    print("\n[5] PARSEANDO ESTADOS DE CUENTA (ZIP)...")
    banco_zip_file = PAQUETE / "ESTADOS_DE_CTA_RENOMBRADO.zip"
    banco_parser = BancoParser()
    movimientos = []
    errores_banco = []
    control_rows = []

    _SYSTEM_DIRS = {"__MACOSX", ".DS_Store", "Thumbs.db", "__pycache__"}

    def _es_ruta_segura(nombre: str) -> bool:
        parts = Path(nombre).parts
        return not any(
            p.startswith(".") or p.startswith("__") or p in _SYSTEM_DIRS
            for p in parts
        )

    with tempfile.TemporaryDirectory() as tmpdir:
        zip_path = os.path.join(tmpdir, "estados.zip")
        with open(zip_path, "wb") as f:
            f.write(banco_zip_file.read_bytes())

        with zipfile.ZipFile(zip_path) as zf:
            compatible = [
                n for n in zf.namelist()
                if _es_ruta_segura(n)
                and not n.startswith("__")
                and not n.endswith("/")
                and any(n.lower().endswith(ext) for ext in (".pdf", ".xlsx", ".xls"))
            ]

            print(f"    ARCHIVOS_COMPATIBLES: {len(compatible)}")

            for name in compatible:
                data = zf.read(name)
                fname = Path(name).name
                ext = Path(name).suffix.lower()
                empresa = Path(name).parts[0] if len(Path(name).parts) > 1 else ""
                error_archivo = ""
                movs_count = 0
                parser_tipo = ""
                moneda_detectada = ""

                try:
                    if ext == ".pdf":
                        from parsers.pdf_parser import BancoPDFParser
                        parser_tipo = "BancoPDFParser"
                        df_pdf, val = BancoPDFParser().parsear_pdf(data, fname)
                        if df_pdf is not None and not df_pdf.empty:
                            movs = banco_parser.parsear_dataframe(df_pdf)
                            movimientos.extend(movs)
                            movs_count = len(movs)
                            if not df_pdf.empty:
                                moneda_detectada = str(df_pdf.iloc[0].get("Moneda", ""))
                        else:
                            errores_banco.append(f"{fname}: sin movimientos")
                    else:
                        parser_tipo = "BancoParser"
                        df = pd.read_excel(io.BytesIO(data))
                        movs = banco_parser.parsear_dataframe(df)
                        movimientos.extend(movs)
                        movs_count = len(movs)
                        if not df.empty:
                            mon_col = next((c for c in df.columns if "moneda" in c.lower()), None)
                            if mon_col:
                                moneda_detectada = str(df.iloc[0].get(mon_col, ""))
                except Exception as exc:
                    error_archivo = str(exc)
                    errores_banco.append(f"{fname}: {exc}")

                status = "OK" if not error_archivo else "ERROR"
                control_rows.append({
                    "ARCHIVO": fname,
                    "EMPRESA": empresa,
                    "TIPO_DOCUMENTO": ext.upper().lstrip("."),
                    "PARSER": parser_tipo,
                    "MOVIMIENTOS_EXTRAIDOS": movs_count,
                    "MONEDA_DETECTADA": moneda_detectada or "NO_DETERMINADA",
                    "ERROR": error_archivo,
                    "ESTATUS": status,
                })
                icon = "+" if status == "OK" else "x"
                print(f"    [{icon}] {fname}: {movs_count} movimientos, {moneda_detectada or 'N/A'}")

    control_df = pd.DataFrame(control_rows)
    print(f"\n    MOVIMIENTOS_BANCARIOS: {len(movimientos)}")
    print(f"    ARCHIVOS_BANCARIOS_PROCESADOS: {len(control_df)}")
    n_ok = len(control_df[control_df["ESTATUS"] == "OK"])
    n_err = len(control_df[control_df["ESTATUS"] == "ERROR"])
    n_con_movs = len(control_df[control_df["MOVIMIENTOS_EXTRAIDOS"] > 0])
    n_sin_movs = len(control_df[control_df["MOVIMIENTOS_EXTRAIDOS"] == 0])
    print(f"    CON_MOVIMIENTOS: {n_con_movs}")
    print(f"    SIN_MOVIMIENTOS: {n_sin_movs}")
    print(f"    CON_ERROR: {n_err}")

    # ── 6. Conciliar egresos ──
    print("\n[6] CONCILIACION EGRESOS...")
    tolerancia = Decimal("1.00")
    dias_tolerancia = 5
    conciliador = Conciliador(tolerancia=tolerancia, dias_tolerancia=dias_tolerancia)
    resultado = conciliador.conciliar(cedula, cfdis, movimientos)
    print(f"    TOTAL_REGISTROS: {resultado.total_registros}")
    print(f"    CONCILIADOS: {resultado.conciliados}")
    print(f"    DIFERENCIAS: {resultado.diferencias}")
    print(f"    SIN_XML: {resultado.sin_xml}")
    print(f"    SIN_BANCO: {resultado.sin_banco}")

    registros_df = registros_a_dataframe(resultado)
    n_propuestas = len(registros_df[registros_df["Estatus"].astype(str).str.contains("PROPUESTA", na=False)])
    n_ambiguos = len(registros_df[registros_df["Estatus"].astype(str).str.contains("AMBIGUO", na=False)])
    n_sin_candidato = len(registros_df[registros_df["Estatus"].astype(str).str.contains("SIN CANDIDATO|SIN_MOVIMIENTO", na=False)])
    print(f"    PROPUESTAS: {n_propuestas}")
    print(f"    AMBIGUOS: {n_ambiguos}")
    print(f"    SIN_CANDIDATO: {n_sin_candidato}")

    # ── 7. Conciliar ingresos ──
    print("\n[7] CONCILIACION INGRESOS...")
    conc_ing = ConciliadorIngresos(tolerancia=tolerancia, dias_tolerancia=dias_tolerancia)
    resultado_ingresos = conc_ing.conciliar(cedula_ingresos, cfdis, movimientos)
    folio = 300
    for reg in resultado_ingresos.registros:
        if reg.cfdi is not None and reg.movimiento is not None and reg.folio_conciliacion is None:
            reg.folio_conciliacion = folio
            folio += 1
    print(f"    TOTAL_REGISTROS_INGRESOS: {resultado_ingresos.total_registros}")
    print(f"    CONCILIADOS_INGRESOS: {resultado_ingresos.conciliados}")
    print(f"    SIN_XML_INGRESOS: {resultado_ingresos.sin_xml}")
    print(f"    SIN_BANCO_INGRESOS: {resultado_ingresos.sin_banco}")
    print(f"    DIFERENCIAS_INGRESOS: {resultado_ingresos.diferencias}")

    # ── 8. Generate Excel ──
    print("\n[8] GENERANDO EXCEL...")
    datos_excel = {
        "RESUMEN": resumen_a_dataframe(resultado),
        "CEDULA_SAT": registros_a_dataframe(resultado),
        "CEDULA_INGRESOS": registros_ingresos_a_dataframe(resultado_ingresos),
        "DETALLE_XML": pd.DataFrame([
            {
                "UUID": c.uuid, "RFC emisor": c.rfc_emisor, "RFC receptor": c.rfc_receptor,
                "Fecha": c.fecha, "Subtotal": float(c.subtotal), "IVA": float(c.iva),
                "Total": float(c.total), "Moneda": c.moneda, "Origen": c.origen_xml,
                "Archivo": c.archivo,
            }
            for c in cfdis
        ]),
        "ESTADOS_CUENTA": pd.DataFrame([
            {
                "Banco": m.banco, "Cuenta": m.cuenta, "Fecha": m.fecha,
                "Concepto": m.concepto, "Cargo": float(m.cargo), "Abono": float(m.abono),
                "Monto": float(m.monto), "Referencia": m.referencia, "Moneda": m.moneda,
            }
            for m in movimientos
        ]),
        "VALIDACION": pd.DataFrame({"Errores": errores_banco}),
    }
    excel_bytes = ExcelGenerator().generar_excel_bytes(datos_excel)

    excel_path = Path(__file__).parent.parent / "outputs" / "evidence_e2e.xlsx"
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    with open(excel_path, "wb") as f:
        f.write(excel_bytes)

    excel_hash = sha256_bytes(excel_bytes)
    excel_size = len(excel_bytes)

    wb_check = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb_check.sheetnames

    print(f"    NOMBRE: {excel_path.name}")
    print(f"    TAMANO: {excel_size:,} bytes")
    print(f"    SHA-256: {excel_hash}")
    print(f"    HOJAS: {sheet_names}")
    print(f"    FECHA_GENERACION: {time.strftime('%Y-%m-%d %H:%M:%S')}")

    ced_sat_cols = list(datos_excel["CEDULA_SAT"].columns)
    print(f"    CEDULA_SAT columnas: {ced_sat_cols}")
    for col in ["CRUCE_ID", "CANDIDATO_CRUCE_ID", "MOVIMIENTO_ID", "CANDIDATO_MOVIMIENTO_ID", "ARCHIVO_ORIGEN", "FILA_ORIGEN"]:
        print(f"    {col}: {'OK' if col in ced_sat_cols else 'FALTA'}")

    expected_sheets = {"RESUMEN", "CEDULA_SAT", "CEDULA_INGRESOS", "DETALLE_XML", "ESTADOS_CUENTA", "VALIDACION"}
    print(f"    HOJAS_ESPERADAS: {sorted(expected_sheets)}")
    print(f"    HOJAS_PRESENTES: {expected_sheets.issubset(set(sheet_names))}")

    elapsed = time.time() - t0
    print(f"\n    E2E completado en {elapsed:.1f}s")
    print("=" * 80)


if __name__ == "__main__":
    main()

"""Genera el Excel de salida con 3 hojas: EGRESOS, INGRESOS, MOVIMIENTOS_BANCARIOS."""
from __future__ import annotations

import hashlib
import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd

from cruce_simple.models import (
    EstatusCruce,
    FilaLayout,
    GrupoPoliza,
    MovimientoBanco,
    ResultadoCruce,
)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def generar_excel(
    filas_egresos: list[FilaLayout],
    filas_ingresos: list[FilaLayout],
    movimientos: list[MovimientoBanco],
    cruce_map: dict[str, str],
    resultados_egresos: list[ResultadoCruce],
    resultados_ingresos: list[ResultadoCruce],
    archivo_salida: str,
) -> dict:
    wb = openpyxl.Workbook()

    ws_eg = wb.active
    ws_eg.title = "EGRESOS"
    ws_eg.append([
        "EMPRESA", "POLIZA", "UUID", "RFC_PROVEEDOR", "PROVEEDOR",
        "FECHA_FACTURA", "FECHA_PAGO", "MONEDA", "TOTAL PESOS", "TOTAL DLS",
        "BANCO", "CRUCE BANCARIO",
    ])
    for f in filas_egresos:
        grupo_key = f"{f.empresa}|{f.poliza}"
        cruce = cruce_map.get(grupo_key, "")
        ws_eg.append([
            f.empresa, f.poliza, f.uuid, f.rfc, f.nombre,
            f.fecha, f.fecha, f.moneda,
            str(f.total_pesos), str(f.total_dls),
            f.banco, cruce,
        ])

    ws_ing = wb.create_sheet("INGRESOS")
    ws_ing.append([
        "EMPRESA", "POLIZA", "UUID", "RFC_CLIENTE", "CLIENTE",
        "FECHA_FACTURA", "FECHA_COBRO", "MONEDA", "TOTAL PESOS", "TOTAL DLS",
        "BANCO", "CRUCE BANCARIO",
    ])
    for f in filas_ingresos:
        grupo_key = f"{f.empresa}|{f.poliza}"
        cruce = cruce_map.get(grupo_key, "")
        ws_ing.append([
            f.empresa, f.poliza, f.uuid, f.rfc, f.nombre,
            f.fecha, f.fecha, f.moneda,
            str(f.total_pesos), str(f.total_dls),
            f.banco, cruce,
        ])

    ws_mov = wb.create_sheet("MOVIMIENTOS_BANCARIOS")
    ws_mov.append([
        "ARCHIVO", "EMPRESA", "BANCO", "CUENTA", "MONEDA",
        "FECHA", "CONCEPTO", "CARGO", "ABONO", "REFERENCIA",
        "CRUCE CARGO", "CRUCE ABONO", "POLIZA RELACIONADA",
    ])
    for m in movimientos:
        ws_mov.append([
            m.archivo, m.empresa, m.banco, m.cuenta, m.moneda,
            m.fecha, m.concepto, str(m.cargo), str(m.abono), m.referencia,
            m.cruce_cargo, m.cruce_abono, m.poliza_relacionada,
        ])

    Path(archivo_salida).parent.mkdir(parents=True, exist_ok=True)
    wb.save(archivo_salida)

    raw = Path(archivo_salida).read_bytes()
    return {
        "archivo": archivo_salida,
        "tamaño": len(raw),
        "sha256": _sha256(raw),
        "hojas": wb.sheetnames,
    }
